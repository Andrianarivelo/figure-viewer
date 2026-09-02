"""Export backend tests: registry contract, content fidelity, cancel safety.

Everything runs headless against a PDF generated on the fly in a temp dir;
no user file is ever read or written. The fixture carries the three things
the exporters must preserve: styled text (a 24 pt heading over 11 pt body),
a ruled 3 x 3 table with cell text, and an embedded PNG image. As in
``test_pdfedit.py``, the source file's hash is asserted unchanged after
every export: exporters read a bytes snapshot and may never touch the disk
copy.

Cancel semantics under test (the shipped contract):

* single-file exporters (text/html/markdown/xlsx/pptx) check the event
  before each page and raise :class:`pdfexport.ExportCancelled` without
  writing any output file (markdown also writes no assets folder);
* ``export_docx`` honours cancel only before pdf2docx starts (the library
  exposes no per-page callback) and raises the same exception;
* ``export_images`` keeps the page files already written and returns a
  partial report whose detail says "cancelled after N of M pages".
"""
from __future__ import annotations

import hashlib
import io
import tempfile
import threading
import unittest
from pathlib import Path

from PIL import Image

from figureviewer import docs as D
from figureviewer import pdfexport as X

TABLE_CELLS = [["Neuron", "Rate", "CV"],
               ["n1", "5.2", "0.81"],
               ["n2", "12.7", "0.45"]]
HEADING = "Neural Data Report"
BODY_LINE = "Spike trains were recorded from CA1."
PAGE2_LINE = "Second page about local field potentials."


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _fixture_bytes() -> bytes:
    """Two-page PDF: heading + body text, ruled 3x3 table, embedded PNG."""
    doc = D.fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 80), HEADING, fontsize=24)
    page.insert_text((72, 120), BODY_LINE, fontsize=11)
    page.insert_text((72, 140), "All rates are given in hertz.", fontsize=11)

    # Ruled grid so find_tables() detects the table with the default
    # lines-based strategy; text is placed inside each cell.
    x0, y0, col_w, row_h = 72, 400, 110, 28
    for r in range(4):
        page.draw_line(D.fitz.Point(x0, y0 + r * row_h),
                       D.fitz.Point(x0 + 3 * col_w, y0 + r * row_h))
    for c in range(4):
        page.draw_line(D.fitz.Point(x0 + c * col_w, y0),
                       D.fitz.Point(x0 + c * col_w, y0 + 3 * row_h))
    for r, row in enumerate(TABLE_CELLS):
        for c, value in enumerate(row):
            page.insert_text((x0 + c * col_w + 6, y0 + r * row_h + 18),
                             value, fontsize=10)

    buffer = io.BytesIO()
    Image.new("RGB", (60, 40), (200, 30, 30)).save(buffer, format="PNG")
    page.insert_image(D.fitz.Rect(430, 100, 520, 160), stream=buffer.getvalue())

    doc.new_page(width=595, height=842).insert_text((72, 100), PAGE2_LINE,
                                                    fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


class _ExportFixtureCase(unittest.TestCase):
    """Shared fixture: the PDF on disk, its bytes, and its integrity hash."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name)
        self.pdf = self.root / "fixture.pdf"
        self.pdf.write_bytes(_fixture_bytes())
        self.data = self.pdf.read_bytes()
        self.source_hash = _sha(self.pdf)

    def _assert_source_untouched(self) -> None:
        self.assertEqual(_sha(self.pdf), self.source_hash,
                         "the raw file on disk must never change")


@unittest.skipUnless(D.HAVE_FITZ, "PyMuPDF is required for export tests")
class ExportRegistryTests(_ExportFixtureCase):
    def test_registry_and_exporters_cover_the_same_formats(self) -> None:
        self.assertEqual(set(X.FORMATS), set(X.EXPORTERS))
        expected_suffixes = {"text": ".txt", "html": ".html", "markdown": ".md",
                             "docx": ".docx", "xlsx": ".xlsx", "pptx": ".pptx",
                             "images": ""}
        self.assertEqual({name: spec.suffix for name, spec in X.FORMATS.items()},
                         expected_suffixes)

    def test_probes_return_available_flag_and_pip_hint(self) -> None:
        for name, spec in X.FORMATS.items():
            with self.subTest(format=name):
                available, hint = spec.probe()
                self.assertIsInstance(available, bool)
                self.assertIsInstance(hint, str)
                if available:
                    self.assertEqual(hint, "")
                else:
                    self.assertIn("pip install", hint,
                                  "an unavailable format must name the fix")

    def test_unavailable_dependency_message_names_the_package(self) -> None:
        # The probe helper itself, independent of what happens to be
        # installed: a module that cannot exist must yield the pip command.
        available, hint = X._module_probe("figview_no_such_module",
                                          "figview-no-such-package",
                                          "Nonexistent export")
        self.assertFalse(available)
        self.assertIn("pip install figview-no-such-package", hint)


@unittest.skipUnless(D.HAVE_FITZ, "PyMuPDF is required for export tests")
class TextHtmlMarkdownExportTests(_ExportFixtureCase):
    def test_text_export_contains_both_pages_in_order(self) -> None:
        ticks: list[tuple[int, int]] = []
        out = self.root / "out.txt"
        report = X.export_text(self.data, out,
                               progress=lambda done, total: ticks.append((done, total)))
        body = out.read_text(encoding="utf-8")
        self.assertEqual(report.pages, 2)
        self.assertIn(HEADING, body)
        self.assertIn(PAGE2_LINE, body)
        self.assertLess(body.index(HEADING), body.index(PAGE2_LINE),
                        "pages must appear in document order")
        self.assertEqual(ticks, [(1, 2), (2, 2)],
                         "progress must tick once per page with the total")
        self._assert_source_untouched()

    def test_page_selection_is_validated_and_honoured(self) -> None:
        out = self.root / "second.txt"
        report = X.export_text(self.data, out, pages=[1])
        body = out.read_text(encoding="utf-8")
        self.assertEqual(report.pages, 1)
        self.assertIn(PAGE2_LINE, body)
        self.assertNotIn(HEADING, body)
        with self.assertRaises(X.ExportError):
            X.export_text(self.data, self.root / "bad.txt", pages=[5])
        with self.assertRaises(X.ExportError):
            X.export_text(self.data, self.root / "none.txt", pages=[])

    def test_html_export_is_self_contained_with_embedded_image(self) -> None:
        out = self.root / "out.html"
        report = X.export_html(self.data, out)
        html = out.read_text(encoding="utf-8")
        self.assertEqual(report.pages, 2)
        self.assertIn(HEADING, html)
        self.assertIn(PAGE2_LINE, html)
        self.assertIn("data:image", html,
                      "images must ride inside the file as data URIs")
        self.assertIn('id="page-1"', html)
        self.assertIn('id="page-2"', html)
        self._assert_source_untouched()

    def test_markdown_headings_tables_and_image_assets(self) -> None:
        out = self.root / "report.md"
        report = X.export_markdown(self.data, out)
        text = out.read_text(encoding="utf-8")
        self.assertEqual(report.pages, 2)
        # The 24 pt line is a heading; the 11 pt body line is not.
        self.assertIn(f"# {HEADING}", text)
        self.assertIn(BODY_LINE, text)
        self.assertNotIn(f"# {BODY_LINE}", text)
        # The ruled grid becomes one pipe table with the fixture cells.
        self.assertIn("| Neuron | Rate | CV |", text)
        self.assertIn("| --- | --- | --- |", text)
        self.assertIn("| n2 | 12.7 | 0.45 |", text)
        # The embedded PNG lands in the assets folder and is referenced
        # relatively so the .md can move together with its folder.
        assets = self.root / "report_assets"
        self.assertTrue(assets.is_dir())
        images = sorted(assets.glob("page1_img*"))
        self.assertEqual(len(images), 1)
        self.assertGreater(images[0].stat().st_size, 0)
        self.assertIn(f"(report_assets/{images[0].name})", text)
        self._assert_source_untouched()

    def test_cancel_after_first_page_writes_no_output(self) -> None:
        for name, exporter in (("text", X.export_text),
                               ("html", X.export_html),
                               ("markdown", X.export_markdown)):
            with self.subTest(format=name):
                cancel = threading.Event()
                out = self.root / f"cancelled.{name}"
                with self.assertRaises(X.ExportCancelled):
                    exporter(self.data, out, cancel=cancel,
                             progress=lambda done, total: cancel.set())
                self.assertFalse(out.exists(),
                                 "a cancelled export must leave no file")
        self.assertFalse((self.root / "cancelled_assets").exists(),
                         "cancelled markdown must not write assets either")
        self._assert_source_untouched()


@unittest.skipUnless(D.HAVE_FITZ, "PyMuPDF is required for export tests")
class ImagesExportTests(_ExportFixtureCase):
    def test_one_file_per_page_with_stable_names(self) -> None:
        out_dir = self.root / "imgs"
        report = X.export_images(self.data, out_dir, dpi=72)
        self.assertEqual(report.pages, 2)
        self.assertEqual([p.name for p in report.files],
                         ["page-001.png", "page-002.png"])
        for path in report.files:
            with Image.open(path) as image:
                self.assertGreater(image.width, 100)
        self._assert_source_untouched()

    def test_cancel_keeps_finished_pages_and_reports_partially(self) -> None:
        cancel = threading.Event()
        out_dir = self.root / "partial"
        report = X.export_images(self.data, out_dir, dpi=72, cancel=cancel,
                                 progress=lambda done, total: cancel.set())
        self.assertEqual(report.pages, 1)
        self.assertEqual([p.name for p in report.files], ["page-001.png"])
        self.assertIn("cancelled after 1 of 2 pages", report.detail)
        self.assertTrue((out_dir / "page-001.png").exists())
        self.assertFalse((out_dir / "page-002.png").exists())

    def test_unsupported_image_format_is_refused(self) -> None:
        with self.assertRaises(X.ExportError):
            X.export_images(self.data, self.root / "bad", image_format="webp")


@unittest.skipUnless(D.HAVE_FITZ and X.xlsx_available()[0],
                     "openpyxl is required for xlsx export tests")
class XlsxExportTests(_ExportFixtureCase):
    def test_detected_table_round_trips_through_openpyxl(self) -> None:
        from openpyxl import load_workbook

        out = self.root / "out.xlsx"
        report = X.export_xlsx(self.data, out)
        self.assertIn("1 table", report.detail)
        workbook = load_workbook(str(out))
        self.assertEqual(workbook.sheetnames, ["Summary", "p1_t1"])
        rows = [[cell.value for cell in row]
                for row in workbook["p1_t1"].iter_rows()]
        self.assertEqual(rows, TABLE_CELLS)
        summary = [[cell.value for cell in row]
                   for row in workbook["Summary"].iter_rows()]
        self.assertEqual(summary[0][:4], ["Sheet", "Page", "Rows", "Columns"])
        self.assertEqual(summary[1][:4], ["p1_t1", 1, 3, 3])
        self._assert_source_untouched()

    def test_zero_tables_raises_instead_of_writing_empty_workbook(self) -> None:
        blank = D.fitz.open()
        blank.new_page(width=200, height=200).insert_text((50, 50), "plain")
        blank_bytes = blank.tobytes()
        blank.close()
        out = self.root / "empty.xlsx"
        with self.assertRaises(X.ExportError) as caught:
            X.export_xlsx(blank_bytes, out)
        self.assertIn("no tables detected", str(caught.exception))
        self.assertFalse(out.exists(),
                         "an honest failure must not leave an empty workbook")


@unittest.skipUnless(D.HAVE_FITZ and X.docx_available()[0],
                     "pdf2docx is required for docx export tests")
class DocxExportTests(_ExportFixtureCase):
    def test_conversion_reloads_with_fixture_text_and_cleans_temp(self) -> None:
        import glob
        import os
        from docx import Document

        pattern = os.path.join(tempfile.gettempdir(), "figview-export-*.pdf")
        before = set(glob.glob(pattern))
        out = self.root / "out.docx"
        with self.assertNoLogs(level="INFO"):
            report = X.export_docx(self.data, out)
        self.assertEqual(report.pages, 2)
        self.assertEqual(set(glob.glob(pattern)), before,
                         "the temp copy fed to pdf2docx must be removed")
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        self.assertIn(HEADING, text)
        self.assertIn(PAGE2_LINE, text)
        self._assert_source_untouched()

    def test_cancel_before_conversion_writes_nothing(self) -> None:
        # pdf2docx exposes no per-page callback, so cancel is honoured only
        # before the conversion starts; a pre-set event must abort cleanly.
        cancel = threading.Event()
        cancel.set()
        out = self.root / "cancelled.docx"
        with self.assertRaises(X.ExportCancelled):
            X.export_docx(self.data, out, cancel=cancel)
        self.assertFalse(out.exists())
        self._assert_source_untouched()


@unittest.skipUnless(D.HAVE_FITZ and X.pptx_available()[0],
                     "python-pptx is required for pptx export tests")
class PptxExportTests(_ExportFixtureCase):
    def test_slides_match_pages_aspect_and_carry_text_in_notes(self) -> None:
        from pptx import Presentation

        out = self.root / "out.pptx"
        report = X.export_pptx(self.data, out, dpi=72)
        self.assertEqual(report.pages, 2)
        deck = Presentation(str(out))
        self.assertEqual(len(deck.slides), 2)
        # Deck aspect equals the fixture page aspect (595 x 842 points).
        self.assertAlmostEqual(deck.slide_width / deck.slide_height,
                               595.0 / 842.0, places=3)
        notes = [slide.notes_slide.notes_text_frame.text
                 for slide in deck.slides]
        self.assertIn(HEADING, notes[0])
        self.assertIn(PAGE2_LINE, notes[1])
        # Each slide is one full-bleed picture of the rendered page.
        for slide in deck.slides:
            picture_shapes = [shape for shape in slide.shapes
                              if shape.shape_type == 13]   # MSO PICTURE
            self.assertEqual(len(picture_shapes), 1)
        self._assert_source_untouched()

    def test_cancel_after_first_page_writes_no_deck(self) -> None:
        cancel = threading.Event()
        out = self.root / "cancelled.pptx"
        with self.assertRaises(X.ExportCancelled):
            X.export_pptx(self.data, out, dpi=72, cancel=cancel,
                          progress=lambda done, total: cancel.set())
        self.assertFalse(out.exists(),
                         "the deck is built in memory; cancel writes nothing")
        self._assert_source_untouched()


if __name__ == "__main__":
    unittest.main()
