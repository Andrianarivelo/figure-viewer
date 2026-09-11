"""Exercise the shipped application using disposable, synthetic files.

Run through the launcher's ``--self-test OUTPUT.json`` switch. The same tests
run from source and inside PyInstaller, where they also detect missing native
libraries and dynamically imported optional export dependencies.
"""
from __future__ import annotations

import json
import platform
import sys
import tempfile
import time
import traceback
from pathlib import Path

WAIT_SECONDS = 4.0
TOTAL_SECONDS = 28.0
IMAGE_COUNT = 16


def run(output: Path) -> int:
    """Write a machine-readable report and return zero only if all checks pass.

    Only generated fixtures and private settings under a temporary directory
    are opened. The user's settings, file associations and files are untouched.
    Callers should enforce a 30-second process timeout as protection against a
    native library hanging, which Python cannot safely interrupt in-process.
    """
    from . import APP_VERSION

    started = time.perf_counter()
    report = {
        "version": APP_VERSION, "frozen": bool(getattr(sys, "frozen", False)),
        "python": platform.python_version(), "platform": platform.platform(),
        "checks": [], "ok": False,
    }
    application = None
    patches = []
    callback_errors = []

    def check(name, operation):
        """Record the result and duration of an independently useful check."""
        tick = time.perf_counter()
        if tick - started > TOTAL_SECONDS:
            raise TimeoutError("Self-test exceeded its total time budget")
        detail = operation()
        report["checks"].append({
            "name": name, "passed": True,
            "seconds": round(time.perf_counter() - tick, 4), "detail": detail,
        })

    def require(condition, message):
        """Use explicit checks so optimized Python still validates results."""
        if not condition:
            raise RuntimeError(message)

    def wait_for(predicate):
        """Pump Tk while waiting for bounded background decode/scan work."""
        deadline = min(started + TOTAL_SECONDS, time.perf_counter() + WAIT_SECONDS)
        while time.perf_counter() < deadline:
            application.root.update()
            require(not callback_errors, "Tk callback failed: " + str(callback_errors))
            if predicate():
                return
            time.sleep(0.01)
        raise TimeoutError("Background operation did not finish in time")

    try:
        from . import app as app_module, ipc
        from PIL import Image
        import fitz

        with tempfile.TemporaryDirectory(prefix="figview-selftest-") as directory:
            workspace = Path(directory)
            # Redirect every persistent application path before creating Tk.
            for module in (app_module, ipc):
                for name, value in (
                    ("SETTINGS_DIR", workspace / "settings"),
                    ("SETTINGS_PATH", workspace / "settings" / "settings.json"),
                    ("INSTANCE_PATH", workspace / "settings" / "instance.json"),
                    ("STARTUP_LOCK_PATH", workspace / "settings" / "startup.lock"),
                ):
                    if hasattr(module, name):
                        patches.append((module, name, getattr(module, name)))
                        setattr(module, name, value)

            images = [workspace / f"image-{index:02d}.png" for index in range(IMAGE_COUNT)]
            for index, path in enumerate(images):
                Image.new("RGB", (96, 64), (20 + index, 80, 160)).save(path)
            pdf = workspace / "document.pdf"
            with fitz.open() as document:
                page = document.new_page(width=288, height=216)
                page.insert_text((25, 40), "FigView packaged release smoke test")
                # A ruled table exercises Excel's actual table extraction path.
                for x in (25, 125, 225):
                    page.draw_line((x, 70), (x, 160))
                for y in (70, 100, 130, 160):
                    page.draw_line((25, y), (225, y))
                for row, values in enumerate((("Item", "Value"), ("Alpha", "10"), ("Beta", "20"))):
                    for column, value in enumerate(values):
                        page.insert_text((30 + 100 * column, 90 + 30 * row), value)
                document.save(pdf)

            def start_gui():
                """Create the actual application and its first viewer window."""
                nonlocal application
                application = app_module.App(select=images[0])
                application.root.report_callback_exception = (
                    lambda kind, value, tb: callback_errors.append(str(value)))
                application.focus_window = lambda window: None
                application.windows[0].withdraw()
                wait_for(lambda: application.windows[0].document is not None)
                require(application.windows[0].document.full_page(0).size == (96, 64),
                        "PNG decoder returned unexpected dimensions")
                return "Tk window created; PNG opened and rendered"

            check("gui_startup_and_png", start_gui)
            window = application.windows[0]

            def open_batches():
                """Rapid successive requests must merge and deduplicate files."""
                application.open_paths(images[:8])
                application.open_paths(images[5:])
                wait_for(lambda: len(window.entries) == IMAGE_COUNT
                         and window.document is not None
                         and window.document.entry.path == images[-1])
                require(len(application.windows) == 1, "Batch opening created extra windows")
                require(len(set(entry.path for entry in window.entries)) == IMAGE_COUNT,
                        "Batch opening introduced duplicate files")
                return {"unique_files": IMAGE_COUNT, "windows": 1}

            check("multiple_file_batches", open_batches)

            def open_via_ipc():
                """Send a real authenticated loopback handoff to the test app."""
                require(application.instance is not None, "IPC server did not start")
                require(ipc.hand_over_to_running_instance(
                    [pdf], instance_path=app_module.INSTANCE_PATH), "IPC handoff rejected")
                wait_for(lambda: window.document is not None
                         and window.document.entry.path == pdf)
                require("smoke test" in window.document.text(0), "PDF text missing")
                require(window.document.full_page(0).width > 0, "PDF render is empty")
                require(len(window.entries) == IMAGE_COUNT + 1, "IPC lost existing files")
                return "Authenticated socket handoff; PDF text and raster rendering verified"

            check("ipc_and_pdf", open_via_ipc)

            def review_and_brand():
                """Verify the shipped review UI, PDF metadata and icon resources."""
                window.text.select_all()
                quote = window.text.selected_text()
                require(bool(quote), "Selectable PDF text is empty")
                uid = window.edit.review.annotate_selection("highlight")
                require(bool(uid), "Direct highlight creation failed")
                session = window.edit.session
                comment = "Review: caf\u00e9, \u03b1 and \u4e2d\u6587"
                session.set_annot_props(0, uid, contents=comment)
                session.undo()
                require(session.annots_on(0)[0].contents == "", "Comment undo failed")
                session.redo()
                window.edit.review.open_panel()
                wait_for(lambda: window.edit.review._job is None
                         and len(window.edit.review.rows) == 1)
                with fitz.open(stream=session.save_bytes(), filetype="pdf") as saved:
                    page = saved[0]
                    annot = next(page.annots())
                    require(annot.info["content"] == comment, "Saved comment changed")
                    require(annot.info["subject"] == quote, "Saved quotation changed")
                with Image.open(application.app_dir / "assets" / "figure-viewer-mark-256.png") as icon:
                    require(icon.size == (256, 256) and icon.mode == "RGBA",
                            "Runtime icon resource is invalid")
                with Image.open(application.app_dir / "figure-viewer.ico") as icon:
                    require((16, 16) in icon.info["sizes"] and (256, 256) in icon.info["sizes"],
                            "Windows icon sizes missing")
                window.edit.review.close_panel()
                # Restore the temporary fixture to a clean session before the
                # remaining checks and normal shutdown, which need no dialogs.
                session.undo()
                session.undo()
                return "Highlight, Unicode comment, undo/redo, review index, PDF save and icons"

            check("pdf_review_and_brand", review_and_brand)

            def video_decode():
                """Generate a tiny AVI and decode through the app video service."""
                import cv2
                import numpy as np
                from .video import VideoSession

                path = workspace / "video.avi"
                writer = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"MJPG"),
                                         10.0, (64, 48))
                require(writer.isOpened(), "MJPEG fixture encoder unavailable")
                try:
                    for index in range(4):
                        writer.write(np.full((48, 64, 3), 40 + index * 20, dtype=np.uint8))
                finally:
                    writer.release()
                session = VideoSession(path, (64, 48), autoplay=False)
                try:
                    wait_for(lambda: session.current_frame() is not None
                             or session.snapshot().state == "error")
                    frame = session.current_frame()
                    require(frame is not None and frame.size == (64, 48),
                            "Video decoder failed: " + str(session.snapshot()))
                    return "MJPEG decoded by bundled OpenCV"
                finally:
                    session.close()

            check("video_decode", video_decode)

            def exports():
                """Validate Office archives by reopening each exported artifact."""
                from . import pdfexport
                from docx import Document
                from openpyxl import load_workbook
                from pptx import Presentation

                data = pdf.read_bytes()
                for extension in ("text", "docx", "xlsx", "pptx"):
                    target = workspace / ("export." + extension)
                    getattr(pdfexport, "export_" + extension)(data, target)
                    require(target.stat().st_size > 0, extension + " export is empty")
                require("smoke test" in (workspace / "export.text").read_text(encoding="utf-8"),
                        "Text export missing content")
                require(len(Document(workspace / "export.docx").paragraphs) > 0,
                        "Word export has no paragraphs")
                workbook = load_workbook(workspace / "export.xlsx")
                try:
                    require(bool(workbook.sheetnames), "Excel export has no sheets")
                finally:
                    workbook.close()
                require(len(Presentation(workspace / "export.pptx").slides) == 1,
                        "PowerPoint export has wrong slide count")
                return ["text", "docx", "xlsx", "pptx"]

            check("pdf_exports", exports)

            def certificate_signing():
                """Sign and verify integrity with a disposable untrusted identity."""
                from . import signing

                password = b"disposable-selftest-password"
                certificate = signing.generate_self_signed(workspace, "Release smoke test", password)
                signed = signing.crypto_sign(pdf, workspace / "signed.pdf", certificate, password)
                results = signing.verify_signatures(signed)
                require(len(results) == 1 and results[0].intact,
                        "Generated PDF signature failed integrity verification")
                return {"intact": results[0].intact, "trusted": results[0].trusted}

            check("certificate_signing", certificate_signing)
            require(not callback_errors, "Tk callback exceptions occurred")
            # Release file handles before TemporaryDirectory removes fixtures.
            application.quit_all()
            application = None
            report["ok"] = True
    except Exception:
        report["error"] = traceback.format_exc()
    finally:
        if application is not None:
            try:
                application.quit_all()
            except Exception:
                report["cleanup_error"] = traceback.format_exc()
        for module, name, value in reversed(patches):
            setattr(module, name, value)
        report["seconds"] = round(time.perf_counter() - started, 4)
        output = Path(output).expanduser().absolute()
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    return 0 if report["ok"] else 1
