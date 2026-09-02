"""PDF export backends: text, HTML, Markdown, DOCX, XLSX, PPTX, page images.

Design rules (enforced by tests in ``tests/test_export.py``)
------------------------------------------------------------
1. **Headless.** This module never imports tkinter; every exporter is a plain
   function testable without a display. All Tk glue (dialogs, worker submit,
   progress marshalling) lives elsewhere.
2. **Snapshot in, files out.** Every exporter takes a ``bytes`` snapshot of
   the PDF, captured by the caller under the document lock, and opens it with
   ``fitz.open(stream=...)``. The live document object is never touched, so
   page operations can never race an export, and the user's raw file is never
   read or written here (except the explicit output path).
3. **Optional dependencies are probed lazily.** pdf2docx, openpyxl and
   python-pptx are looked up with ``importlib.util.find_spec`` in the style
   of :func:`figureviewer.video.backend_available`. A missing package yields
   an :class:`ExportError` naming the exact ``pip install`` command; nothing
   optional is imported at module import time.
4. **Cancellable and observable.** Every exporter checks its ``cancel`` event
   between pages and calls ``progress(done, total)`` after each page. In-file
   exporters abort cleanly (no output file is written) by raising
   :class:`ExportCancelled`; :func:`export_images` keeps the page files that
   already finished and returns a partial, honestly labelled report.
5. **Honest output.** Zero detected tables raises ``ExportError('no tables
   detected')`` instead of writing an empty workbook; scanned pages without a
   text layer fall back to the ``ocr_words`` provider when one is given.
"""
from __future__ import annotations

import html as html_module
import importlib.util
import io
import logging
import os
import tempfile
import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional, Sequence

from . import docs as D

#: Fraction of (character-weighted) span sizes that must sit below a font
#: size before it is considered a heading candidate in Markdown export.
HEADING_PERCENTILE = 0.85
#: A heading must also be clearly larger than the document's median size,
#: otherwise a uniformly-sized document would export as wall-to-wall headings.
HEADING_MEDIAN_RATIO = 1.15
#: Native word count below which a page counts as "no text layer" and the
#: ``ocr_words`` provider is consulted (mirrors EditSession.page_text).
OCR_WORD_THRESHOLD = 3
#: Default render resolution for image-based exports (pptx slides, images).
DEFAULT_EXPORT_DPI = 150

#: Font-name substrings that mark a span as monospace in Markdown export.
_MONO_HINTS = ("courier", "consolas", "mono")

ProgressFn = Optional[Callable[[int, int], None]]
OcrWordsFn = Optional[Callable[[int], list]]


class ExportError(RuntimeError):
    """A user-facing export failure (missing package, bad range, no tables)."""


class ExportCancelled(ExportError):
    """The cancel event was set; no (further) output was written."""


@dataclass(frozen=True)
class ExportReport:
    """What an exporter produced: where, how many pages, and any fine print."""

    out_path: Path
    pages: int
    detail: str = ""
    files: tuple[Path, ...] = ()


@dataclass(frozen=True)
class FormatSpec:
    """One entry of the :data:`FORMATS` registry.

    ``probe`` returns ``(available, hint)``; when unavailable the hint names
    the exact ``pip install`` command (or missing component) so the UI can
    show an actionable message on the disabled row.
    """

    label: str
    suffix: str
    probe: Callable[[], tuple[bool, str]]
    needs_ocr_text: bool


# ---------------------------------------------------------------------------
# Dependency probes (lazy, import-free, video.backend_available style)


_PYMUPDF_HINT = ("PDF export requires PyMuPDF.\n\nInstall it with:\n"
                 "    pip install pymupdf")


def fitz_available() -> tuple[bool, str]:
    """PyMuPDF probe: every export path needs it to read the snapshot."""
    if D.HAVE_FITZ:
        return True, ""
    return False, _PYMUPDF_HINT


def _module_probe(module: str, package: str, label: str) -> tuple[bool, str]:
    ok, hint = fitz_available()
    if not ok:
        return False, hint
    if importlib.util.find_spec(module) is not None:
        return True, ""
    return False, (f"{label} requires the {package} package.\n\n"
                   f"Install it with:\n    pip install {package}")


def docx_available() -> tuple[bool, str]:
    return _module_probe("pdf2docx", "pdf2docx", "Word export")


def xlsx_available() -> tuple[bool, str]:
    return _module_probe("openpyxl", "openpyxl", "Excel export")


def pptx_available() -> tuple[bool, str]:
    return _module_probe("pptx", "python-pptx", "PowerPoint export")


def _require(probe: Callable[[], tuple[bool, str]]) -> None:
    ok, hint = probe()
    if not ok:
        raise ExportError(hint)


# ---------------------------------------------------------------------------
# Shared plumbing


def _open_snapshot(data: bytes):
    """Open the caller's byte snapshot; never a path, never the live doc."""
    _require(fitz_available)
    try:
        return D.fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ExportError(f"could not open the PDF snapshot: {exc}") from exc


def _select_pages(doc, pages: Optional[Sequence[int]]) -> list[int]:
    """Validate and normalise a 0-based page selection."""
    total = doc.page_count
    if pages is None:
        return list(range(total))
    result = sorted({int(p) for p in pages})
    if not result:
        raise ExportError("no pages selected")
    for index in result:
        if index < 0 or index >= total:
            raise ExportError(
                f"page {index + 1} is out of range (document has {total} pages)")
    return result


def _check_cancel(cancel: Optional[threading.Event]) -> None:
    if cancel is not None and cancel.is_set():
        raise ExportCancelled("export cancelled")


def _tick(progress: ProgressFn, done: int, total: int) -> None:
    if progress is not None:
        progress(done, total)


def _word_box_and_text(word) -> tuple[tuple[float, float, float, float], str]:
    """Accept both OcrWord-style objects (rect_pts/text) and fitz word tuples."""
    rect = getattr(word, "rect_pts", None)
    if rect is not None:
        return tuple(rect), str(getattr(word, "text", ""))
    rect = getattr(word, "rect", None)
    if rect is not None:
        text = str(getattr(word, "text", ""))
        return (rect[0], rect[1], rect[2], rect[3]), text
    # fitz "words" tuple: (x0, y0, x1, y1, word, block, line, word_no)
    return (word[0], word[1], word[2], word[3]), str(word[4])


def _words_to_text(words) -> str:
    """Assemble OCR words into reading-order lines by their page geometry."""
    boxed = [_word_box_and_text(w) for w in words or ()]
    boxed = [(box, text) for box, text in boxed if text.strip()]
    if not boxed:
        return ""
    boxed.sort(key=lambda item: (item[0][1], item[0][0]))
    lines: list[list[tuple[tuple[float, float, float, float], str]]] = []
    for box, text in boxed:
        if lines:
            prev_box = lines[-1][-1][0]
            prev_mid = (prev_box[1] + prev_box[3]) / 2.0
            if abs(((box[1] + box[3]) / 2.0) - prev_mid) <= max(
                    2.0, (box[3] - box[1]) * 0.6):
                lines[-1].append((box, text))
                continue
        lines.append([(box, text)])
    rendered = []
    for line in lines:
        line.sort(key=lambda item: item[0][0])
        rendered.append(" ".join(text for _box, text in line))
    return "\n".join(rendered)


def _needs_ocr(page) -> bool:
    return len(page.get_text("words")) < OCR_WORD_THRESHOLD


def _ocr_page_text(index: int, ocr_words: OcrWordsFn) -> str:
    """Best-effort text from the OCR provider; empty on any failure."""
    if ocr_words is None:
        return ""
    try:
        return _words_to_text(ocr_words(index))
    except Exception:
        return ""


def _page_text(page, index: int, ocr_words: OcrWordsFn) -> str:
    """Native text with an OCR fallback for pages without a text layer."""
    native = page.get_text("text")
    if not _needs_ocr(page):
        return native
    return _ocr_page_text(index, ocr_words) or native


# ---------------------------------------------------------------------------
# Plain text


def export_text(data: bytes, out_path: Path, *,
                pages: Optional[list[int]] = None,
                progress: ProgressFn = None,
                cancel: Optional[threading.Event] = None,
                ocr_words: OcrWordsFn = None) -> ExportReport:
    """Write the document's text layer (page by page) to a UTF-8 file."""
    out_path = Path(out_path)
    doc = _open_snapshot(data)
    try:
        page_list = _select_pages(doc, pages)
        total = len(page_list)
        chunks: list[str] = []
        ocr_pages = 0
        for done, index in enumerate(page_list, start=1):
            _check_cancel(cancel)
            page = doc[index]
            text = _page_text(page, index, ocr_words)
            if _needs_ocr(page) and text.strip():
                ocr_pages += 1
            chunks.append(text.rstrip("\n"))
            _tick(progress, done, total)
    finally:
        doc.close()
    out_path.write_text("\n\n".join(chunks) + "\n", encoding="utf-8")
    detail = f"{total} page{'s' if total != 1 else ''} of text"
    if ocr_pages:
        detail += f", {ocr_pages} via OCR"
    return ExportReport(out_path, total, detail)


# ---------------------------------------------------------------------------
# HTML

_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
  :root {{ color-scheme: light; }}
  body {{ margin: 0; padding: 2rem 1rem; background: #eceae6;
         font-family: Georgia, 'Times New Roman', serif; color: #1f1d1a; }}
  main {{ max-width: 52rem; margin: 0 auto; }}
  section.page {{ background: #fffdf9; border: 1px solid #d8d4cc;
                  border-radius: 6px; padding: 2.2rem 2.6rem;
                  margin: 0 0 1.6rem 0; box-shadow: 0 1px 4px rgba(0,0,0,.06);
                  overflow-x: auto; }}
  section.page > header {{ font: 600 .72rem/1 system-ui, sans-serif;
                           letter-spacing: .08em; text-transform: uppercase;
                           color: #8a857b; margin-bottom: 1.1rem; }}
  img {{ max-width: 100%; height: auto; }}
  p {{ line-height: 1.5; }}
  .ocr {{ font-style: italic; }}
</style>
</head>
<body>
<main>
{body}
</main>
</body>
</html>
"""


def export_html(data: bytes, out_path: Path, *,
                pages: Optional[list[int]] = None,
                progress: ProgressFn = None,
                cancel: Optional[threading.Event] = None,
                ocr_words: OcrWordsFn = None) -> ExportReport:
    """Write a single self-contained HTML file (images as data URIs).

    PyMuPDF's ``xhtml`` extraction already embeds every image as a base64
    ``data:`` URI, so the output needs no side-car files and survives being
    mailed around as one document.
    """
    out_path = Path(out_path)
    doc = _open_snapshot(data)
    try:
        page_list = _select_pages(doc, pages)
        total = len(page_list)
        sections: list[str] = []
        for done, index in enumerate(page_list, start=1):
            _check_cancel(cancel)
            page = doc[index]
            body = page.get_text("xhtml")
            if _needs_ocr(page):
                ocr_text = _ocr_page_text(index, ocr_words)
                if ocr_text:
                    paragraphs = "\n".join(
                        f"<p class=\"ocr\">{html_module.escape(line)}</p>"
                        for line in ocr_text.splitlines() if line.strip())
                    body += "\n" + paragraphs
            sections.append(
                f"<section class=\"page\" id=\"page-{index + 1}\">\n"
                f"<header>Page {index + 1}</header>\n{body}\n</section>")
            _tick(progress, done, total)
        title = html_module.escape(out_path.stem)
    finally:
        doc.close()
    out_path.write_text(
        _HTML_TEMPLATE.format(title=title, body="\n".join(sections)),
        encoding="utf-8")
    return ExportReport(out_path, total,
                        f"self-contained HTML, {total} page"
                        f"{'s' if total != 1 else ''}")


# ---------------------------------------------------------------------------
# Markdown


def _weighted_percentile(weighted_sizes: list[tuple[float, int]],
                         fraction: float) -> float:
    """Character-weighted percentile of span font sizes."""
    if not weighted_sizes:
        return 0.0
    ordered = sorted(weighted_sizes)
    total = sum(weight for _size, weight in ordered)
    threshold = total * fraction
    seen = 0
    for size, weight in ordered:
        seen += weight
        if seen >= threshold:
            return size
    return ordered[-1][0]


def _collect_span_sizes(doc, page_list: list[int]) -> list[tuple[float, int]]:
    sizes: list[tuple[float, int]] = []
    for index in page_list:
        for block in doc[index].get_text("dict")["blocks"]:
            if block.get("type") != 0:
                continue
            for line in block.get("lines", ()):
                for span in line.get("spans", ()):
                    text = span.get("text", "").strip()
                    if text:
                        sizes.append((round(float(span.get("size", 0.0)), 1),
                                      len(text)))
    return sizes


def _heading_levels(sizes: list[tuple[float, int]]) -> dict[float, int]:
    """Map heading-worthy font sizes to Markdown levels (1 = ``#``)."""
    if not sizes:
        return {}
    p85 = _weighted_percentile(sizes, HEADING_PERCENTILE)
    median = _weighted_percentile(sizes, 0.50)
    threshold = max(p85, median * HEADING_MEDIAN_RATIO)
    heading_sizes = sorted({size for size, _w in sizes if size >= threshold},
                           reverse=True)
    return {size: min(rank + 1, 3) for rank, size in enumerate(heading_sizes)}


def _md_escape_cell(value) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", " ").strip()


def _pipe_table(rows: list[list]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normal = [[_md_escape_cell(cell) for cell in row] + [""] * (width - len(row))
              for row in rows]
    lines = ["| " + " | ".join(normal[0]) + " |",
             "|" + "|".join([" --- "] * width) + "|"]
    lines.extend("| " + " | ".join(row) + " |" for row in normal[1:])
    return "\n".join(lines)


def _md_span(span) -> str:
    text = span.get("text", "")
    if not text.strip():
        return text
    flags = int(span.get("flags", 0))
    font = str(span.get("font", "")).lower()
    mono = bool(flags & 8) or any(hint in font for hint in _MONO_HINTS)
    if mono:
        return f"`{text.strip()}`"
    bold = bool(flags & 16) or "bold" in font or "black" in font or "heavy" in font
    italic = bool(flags & 2) or "italic" in font or "oblique" in font
    stripped = text.strip()
    lead = text[:len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    if bold and italic:
        stripped = f"***{stripped}***"
    elif bold:
        stripped = f"**{stripped}**"
    elif italic:
        stripped = f"*{stripped}*"
    return f"{lead}{stripped}{trail}"


def _dominant_size(line) -> float:
    best_size, best_weight = 0.0, -1
    for span in line.get("spans", ()):
        weight = len(span.get("text", "").strip())
        if weight > best_weight:
            best_weight = weight
            best_size = round(float(span.get("size", 0.0)), 1)
    return best_size


def _block_in_tables(block_bbox, table_bboxes) -> Optional[int]:
    """Index of the table whose bbox contains the block's centre, else None."""
    cx = (block_bbox[0] + block_bbox[2]) / 2.0
    cy = (block_bbox[1] + block_bbox[3]) / 2.0
    for position, (x0, y0, x1, y1) in enumerate(table_bboxes):
        if x0 <= cx <= x1 and y0 <= cy <= y1:
            return position
    return None


def export_markdown(data: bytes, out_path: Path, *,
                    pages: Optional[list[int]] = None,
                    progress: ProgressFn = None,
                    cancel: Optional[threading.Event] = None,
                    ocr_words: OcrWordsFn = None) -> ExportReport:
    """Write structured Markdown: headings, emphasis, tables, linked images.

    Headings come from a document-level font-size percentile, so 'large for
    this document' rather than any absolute size decides. Images are written
    into an assets folder beside the ``.md`` and referenced relatively; the
    whole export is buffered so a cancel writes nothing at all.
    """
    out_path = Path(out_path)
    doc = _open_snapshot(data)
    try:
        page_list = _select_pages(doc, pages)
        total = len(page_list)
        levels = _heading_levels(_collect_span_sizes(doc, page_list))
        assets_dirname = f"{out_path.stem}_assets"
        assets: list[tuple[str, bytes]] = []   # buffered: cancel writes nothing
        parts: list[str] = []
        table_count = 0
        for done, index in enumerate(page_list, start=1):
            _check_cancel(cancel)
            page = doc[index]
            tables = list(page.find_tables().tables)
            table_bboxes = [tuple(t.bbox) for t in tables]
            emitted = [False] * len(tables)

            def emit_table(position: int) -> None:
                nonlocal table_count
                if emitted[position]:
                    return
                emitted[position] = True
                table_count += 1
                parts.append(_pipe_table(tables[position].extract()))

            image_no = 0
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") == 1:
                    ext = block.get("ext") or "png"
                    payload = block.get("image") or b""
                    if not payload:
                        continue
                    image_no += 1
                    filename = f"page{index + 1}_img{image_no}.{ext}"
                    assets.append((filename, payload))
                    parts.append(f"![Page {index + 1} image {image_no}]"
                                 f"({assets_dirname}/{filename})")
                    continue
                position = _block_in_tables(block.get("bbox", (0, 0, 0, 0)),
                                            table_bboxes)
                if position is not None:
                    emit_table(position)
                    continue
                lines_out: list[str] = []
                for line in block.get("lines", ()):
                    text = "".join(_md_span(s) for s in line.get("spans", ()))
                    text = text.strip()
                    if not text:
                        continue
                    level = levels.get(_dominant_size(line))
                    if level:
                        if lines_out:
                            parts.append("\n".join(lines_out))
                            lines_out = []
                        parts.append(f"{'#' * level} {text}")
                    else:
                        lines_out.append(text)
                if lines_out:
                    parts.append("\n".join(lines_out))
            for position in range(len(tables)):
                emit_table(position)
            if _needs_ocr(page):
                ocr_text = _ocr_page_text(index, ocr_words)
                if ocr_text:
                    parts.extend(line for line in ocr_text.splitlines()
                                 if line.strip())
            _tick(progress, done, total)
    finally:
        doc.close()
    if assets:
        assets_dir = out_path.parent / assets_dirname
        assets_dir.mkdir(parents=True, exist_ok=True)
        for filename, payload in assets:
            (assets_dir / filename).write_bytes(payload)
    out_path.write_text("\n\n".join(parts) + "\n", encoding="utf-8")
    detail = (f"{total} page{'s' if total != 1 else ''}, "
              f"{table_count} table{'s' if table_count != 1 else ''}, "
              f"{len(assets)} image{'s' if len(assets) != 1 else ''}")
    return ExportReport(out_path, total, detail)


# ---------------------------------------------------------------------------
# DOCX (pdf2docx)


def export_docx(data: bytes, out_path: Path, *,
                pages: Optional[list[int]] = None,
                progress: ProgressFn = None,
                cancel: Optional[threading.Event] = None,
                ocr_words: OcrWordsFn = None) -> ExportReport:
    """Convert to Word via pdf2docx (which needs a real input path).

    The snapshot bytes go to a private temp file that is always removed in
    ``finally``. pdf2docx's public ``convert`` exposes no per-page callback,
    so progress is reported per call (0 of 1, then 1 of 1); ``cancel`` is
    honoured before the conversion starts. ``ocr_words`` is accepted for
    signature symmetry: baking OCR text into the temp copy is the caller's
    job (a make-searchable pass) because pdf2docx only reads real glyphs.
    """
    del ocr_words   # see docstring: handled upstream by make_searchable
    _require(docx_available)
    out_path = Path(out_path)
    doc = _open_snapshot(data)
    try:
        page_list = _select_pages(doc, pages)
    finally:
        doc.close()
    _check_cancel(cancel)
    _tick(progress, 0, 1)

    from pdf2docx import Converter

    # pdf2docx narrates every layout decision at INFO through the *root*
    # logger (its package __init__ even calls logging.basicConfig, so the
    # narration reaches a real handler). A per-logger level cannot silence
    # that; scope a global disable to the conversion and restore it after.
    previous_disable = logging.root.manager.disable
    logging.disable(logging.INFO)
    handle, tmp_name = tempfile.mkstemp(suffix=".pdf", prefix="figview-export-")
    converter = None
    try:
        with os.fdopen(handle, "wb") as tmp_file:
            tmp_file.write(data)
        converter = Converter(tmp_name)
        converter.convert(str(out_path), pages=page_list)
    except ExportError:
        raise
    except Exception as exc:
        raise ExportError(f"pdf2docx conversion failed: {exc}") from exc
    finally:
        if converter is not None:
            try:
                converter.close()
            except Exception:
                pass
        logging.disable(previous_disable)
        try:
            os.remove(tmp_name)
        except OSError:
            pass
    _tick(progress, 1, 1)
    return ExportReport(out_path, len(page_list), "converted with pdf2docx")


# ---------------------------------------------------------------------------
# XLSX (openpyxl)


def export_xlsx(data: bytes, out_path: Path, *,
                pages: Optional[list[int]] = None,
                progress: ProgressFn = None,
                cancel: Optional[threading.Event] = None,
                ocr_words: OcrWordsFn = None) -> ExportReport:
    """One worksheet per detected table plus a summary sheet.

    Zero detected tables raises ``ExportError('no tables detected')`` so the
    UI shows an honest message instead of an empty workbook. The workbook is
    built in memory and only saved at the end: a cancel writes nothing.
    """
    del ocr_words   # tables need ruled lines and glyph geometry, not OCR text
    _require(xlsx_available)
    out_path = Path(out_path)

    from openpyxl import Workbook

    doc = _open_snapshot(data)
    try:
        page_list = _select_pages(doc, pages)
        total = len(page_list)
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Sheet", "Page", "Rows", "Columns", "Bounding box (pt)"])
        table_count = 0
        for done, index in enumerate(page_list, start=1):
            _check_cancel(cancel)
            page = doc[index]
            for position, table in enumerate(page.find_tables().tables, start=1):
                table_count += 1
                name = f"p{index + 1}_t{position}"
                sheet = workbook.create_sheet(title=name)
                for row in table.extract():
                    sheet.append(["" if cell is None else cell for cell in row])
                x0, y0, x1, y1 = table.bbox
                summary.append([name, index + 1, table.row_count,
                                table.col_count,
                                f"({x0:.0f}, {y0:.0f}, {x1:.0f}, {y1:.0f})"])
            _tick(progress, done, total)
    finally:
        doc.close()
    if table_count == 0:
        raise ExportError("no tables detected")
    workbook.save(str(out_path))
    return ExportReport(out_path, total,
                        f"{table_count} table{'s' if table_count != 1 else ''} "
                        f"from {total} page{'s' if total != 1 else ''}")


# ---------------------------------------------------------------------------
# PPTX (python-pptx)


def export_pptx(data: bytes, out_path: Path, *,
                pages: Optional[list[int]] = None,
                progress: ProgressFn = None,
                cancel: Optional[threading.Event] = None,
                ocr_words: OcrWordsFn = None,
                dpi: int = DEFAULT_EXPORT_DPI) -> ExportReport:
    """Image-based slides with the page text in each slide's notes.

    The deck is sized to the first exported page's aspect ratio; every page
    becomes one full-bleed PNG rendered at ``dpi``. Slide notes carry the
    page's text (native, or OCR via ``ocr_words``), which is what makes the
    deck searchable. Built in memory; a cancel writes nothing.
    """
    _require(pptx_available)
    out_path = Path(out_path)

    from pptx import Presentation
    from pptx.util import Emu

    emu_per_point = 12700
    doc = _open_snapshot(data)
    try:
        page_list = _select_pages(doc, pages)
        total = len(page_list)
        first_rect = doc[page_list[0]].rect
        deck = Presentation()
        deck.slide_width = Emu(int(first_rect.width * emu_per_point))
        deck.slide_height = Emu(int(first_rect.height * emu_per_point))
        blank_layout = deck.slide_layouts[6]
        for done, index in enumerate(page_list, start=1):
            _check_cancel(cancel)
            page = doc[index]
            png = page.get_pixmap(dpi=int(dpi)).tobytes("png")
            slide = deck.slides.add_slide(blank_layout)
            slide.shapes.add_picture(io.BytesIO(png), 0, 0,
                                     width=deck.slide_width,
                                     height=deck.slide_height)
            notes = _page_text(page, index, ocr_words).strip()
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
            _tick(progress, done, total)
    finally:
        doc.close()
    deck.save(str(out_path))
    return ExportReport(out_path, total,
                        "image-based slides with searchable notes")


# ---------------------------------------------------------------------------
# Page images


def export_images(data: bytes, out_path: Path, *,
                  pages: Optional[list[int]] = None,
                  progress: ProgressFn = None,
                  cancel: Optional[threading.Event] = None,
                  ocr_words: OcrWordsFn = None,
                  dpi: int = DEFAULT_EXPORT_DPI,
                  image_format: str = "png",
                  stem: str = "page") -> ExportReport:
    """Render each page into ``out_path`` (a folder) as PNG or JPEG.

    Unlike the single-file exporters, pages that already finished are real
    files on disk when a cancel arrives; deleting the user's finished output
    would be worse than keeping it, so this exporter returns a partial report
    whose detail says exactly how far it got.
    """
    del ocr_words   # a rendered page already shows everything OCR would add
    _require(fitz_available)
    fmt = image_format.lower().lstrip(".")
    if fmt in ("jpg", "jpeg"):
        fmt, suffix = "jpeg", ".jpg"
    elif fmt == "png":
        suffix = ".png"
    else:
        raise ExportError(f"unsupported image format {image_format!r}; "
                          "use 'png' or 'jpeg'")
    out_dir = Path(out_path)
    doc = _open_snapshot(data)
    try:
        page_list = _select_pages(doc, pages)
        total = len(page_list)
        out_dir.mkdir(parents=True, exist_ok=True)
        files: list[Path] = []
        cancelled = False
        for done, index in enumerate(page_list, start=1):
            if cancel is not None and cancel.is_set():
                cancelled = True
                break
            target = out_dir / f"{stem}-{index + 1:03d}{suffix}"
            doc[index].get_pixmap(dpi=int(dpi)).save(str(target))
            files.append(target)
            _tick(progress, done, total)
    finally:
        doc.close()
    if cancelled:
        detail = f"cancelled after {len(files)} of {total} pages"
    else:
        detail = (f"{len(files)} {fmt.upper()} image"
                  f"{'s' if len(files) != 1 else ''} at {int(dpi)} dpi")
    return ExportReport(out_dir, len(files), detail, tuple(files))


# ---------------------------------------------------------------------------
# Registry


#: Ordered registry the export dialog renders top to bottom. ``suffix`` is
#: the default file suffix ("" for the folder-based image export).
FORMATS: dict[str, FormatSpec] = {
    "text": FormatSpec("Plain text", ".txt", fitz_available, True),
    "html": FormatSpec("HTML (self-contained)", ".html", fitz_available, True),
    "markdown": FormatSpec("Markdown", ".md", fitz_available, True),
    "docx": FormatSpec("Word document", ".docx", docx_available, True),
    "xlsx": FormatSpec("Excel tables", ".xlsx", xlsx_available, False),
    "pptx": FormatSpec("PowerPoint slides", ".pptx", pptx_available, True),
    "images": FormatSpec("Page images", "", fitz_available, False),
}

#: Exporter function for each registry name, same keyword contract everywhere.
EXPORTERS: dict[str, Callable[..., ExportReport]] = {
    "text": export_text,
    "html": export_html,
    "markdown": export_markdown,
    "docx": export_docx,
    "xlsx": export_xlsx,
    "pptx": export_pptx,
    "images": export_images,
}
